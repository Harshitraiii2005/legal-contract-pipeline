"""Build a redlined .docx with tracked-change styling."""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.util import Inches

from app.graph.state import ClauseExtract, RedlineEdit
from app.core.logging import get_logger

logger = get_logger(__name__)

RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x70, 0x00)
GREY = RGBColor(0x60, 0x60, 0x60)


def build_redlined_docx(
    contract_name: str,
    clauses: list[ClauseExtract],
    edits: list[RedlineEdit],
    overall_score: int,
) -> bytes:
    """Return bytes of a .docx showing original vs revised clauses."""
    doc = Document()
    _set_page_margins(doc)

    # Title
    title = doc.add_heading(f"Redlined Contract: {contract_name}", level=0)
    title.runs[0].font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    doc.add_paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}  |  "
        f"Overall Risk Score: {overall_score}/100"
    ).runs[0].font.color.rgb = GREY

    doc.add_paragraph()  # spacer

    # Index edits by clause_id
    edit_map = {e.clause_id: e for e in edits}

    for clause in clauses:
        edit = edit_map.get(clause.id)

        # Heading
        h = doc.add_heading(
            f"{clause.id}. {clause.heading or clause.type.replace('_', ' ').title()}",
            level=2,
        )
        h.runs[0].font.color.rgb = RED if edit else RGBColor(0x1A, 0x1A, 0x2E)

        if edit:
            # Original (strikethrough, red)
            p_orig = doc.add_paragraph()
            r = p_orig.add_run("ORIGINAL: ")
            r.bold = True
            r.font.color.rgb = RED
            r_text = p_orig.add_run(clause.text)
            r_text.font.color.rgb = RED
            _apply_strikethrough(r_text)

            doc.add_paragraph()

            # Revised (green)
            p_rev = doc.add_paragraph()
            r2 = p_rev.add_run("REVISED: ")
            r2.bold = True
            r2.font.color.rgb = GREEN
            r_text2 = p_rev.add_run(edit.revised_text)
            r_text2.font.color.rgb = GREEN

            # Attorney note
            if edit.attorney_note:
                note = doc.add_paragraph()
                nr = note.add_run(f"⚖ Attorney note: {edit.attorney_note}")
                nr.italic = True
                nr.font.size = Pt(9)
                nr.font.color.rgb = GREY
        else:
            doc.add_paragraph(clause.text)

        doc.add_paragraph()  # spacer between clauses

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Helpers ───────────────────────────────────────────────────────────────────

def _apply_strikethrough(run) -> None:
    rPr = run._r.get_or_add_rPr()
    strike = OxmlElement("w:strike")
    strike.set(qn("w:val"), "true")
    rPr.append(strike)


def _set_page_margins(doc: Document, margin_inches: float = 1.0) -> None:
    section = doc.sections[0]
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Inches(margin_inches))
